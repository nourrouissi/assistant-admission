"""
Import du fichier Excel des formations vers la base de données (le « L » de l'ETL).

Séparé de la commande Django (management/commands/import_formations.py) pour
pouvoir être appelé et testé directement, sans passer par la ligne de commande.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from django.db import transaction
from openpyxl import load_workbook

from .models import Formation
from .normalization import (
    DiplomaType,
    clean_spaces,
    clean_specialty,
    consistency_warnings,
    fold,
    normalize_diploma_type,
    parse_bac_rule,
    parse_level_requirement,
    parse_nationality,
    school_city,
)

# Colonnes attendues, reconnues par leur en-tête « plié » (sans accents ni espaces
# parasites) et non par leur position : l'import résiste à un changement d'ordre
# des colonnes dans le fichier.
EXPECTED_COLUMNS = {
    "diplome": "diploma",
    "specialite": "specialty",
    "ecole": "school",
    "niveau eligible": "level",
    "nationalite eligible": "nationality",
    "nature du bac": "bac",
}


@dataclass
class ImportReport:
    """Bilan d'un import, affiché par la commande."""

    created: int = 0
    updated: int = 0
    skipped: int = 0
    warnings: list[str] = field(default_factory=list)

    @property
    def total(self) -> int:
        return self.created + self.updated


def _map_header(header_row: tuple) -> dict[str, int]:
    """Associe chaque champ attendu à l'index de sa colonne ; erreur claire si une colonne manque."""
    mapping: dict[str, int] = {}
    for index, title in enumerate(header_row):
        key = EXPECTED_COLUMNS.get(fold(str(title or "")))
        if key:
            mapping[key] = index
    missing = set(EXPECTED_COLUMNS.values()) - set(mapping)
    if missing:
        raise ValueError(f"Colonnes manquantes dans le fichier Excel : {', '.join(sorted(missing))}")
    return mapping


def build_formation_fields(row: dict[str, str]) -> dict:
    """
    Transforme une ligne brute (colonne -> texte) en champs du modèle Formation.
    Toutes les règles de normalisation sont appliquées ici ; les valeurs
    d'origine sont conservées à côté des valeurs normalisées.
    """
    diploma_label = clean_spaces(row["diploma"])
    diploma_type = normalize_diploma_type(diploma_label)
    specialty = clean_specialty(row["specialty"])
    school = clean_spaces(row["school"]).upper()
    level = parse_level_requirement(row["level"])
    bac_rule, bac_sections = parse_bac_rule(row["bac"])
    type_label = DiplomaType.LABELS.get(diploma_type or "", "")

    return {
        "diploma_label": diploma_label,
        "diploma_type": diploma_type,
        "specialty": specialty,
        "school": school,
        "city": school_city(school),
        "level_raw": clean_spaces(row["level"]),
        "min_level": level.min_level,
        "level_alternatives": level.alternatives,
        "alternatives_same_sector": level.alternatives_same_sector,
        "nationality_raw": clean_spaces(row["nationality"]),
        "nationality_required": parse_nationality(row["nationality"]),
        "bac_raw": clean_spaces(row["bac"]),
        "bac_rule": bac_rule,
        "bac_sections": bac_sections,
        # Texte « plié » qui regroupe tout ce qu'on peut chercher
        "search_text": fold(f"{diploma_label} {type_label} {specialty} {school}"),
        "import_warnings": consistency_warnings(diploma_type, level, bac_rule),
    }


@transaction.atomic
def import_formations(path: Path) -> ImportReport:
    """
    Importe (ou met à jour) toutes les formations du fichier Excel.

    - Idempotent : relancer l'import ne crée pas de doublons. Chaque ligne est
      retrouvée par sa clé naturelle (diplôme + spécialité + école), la même
      que la contrainte d'unicité du modèle.
    - Atomique (@transaction.atomic) : en cas d'erreur, rien n'est enregistré.
    """
    report = ImportReport()
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    columns = _map_header(next(rows))

    for line_number, values in enumerate(rows, start=2):
        row = {key: str(values[idx] or "") for key, idx in columns.items()}
        if not clean_spaces(row["diploma"]) or not clean_spaces(row["specialty"]):
            report.skipped += 1  # ligne vide en fin de fichier
            continue

        fields = build_formation_fields(row)
        natural_key = {
            "diploma_label": fields.pop("diploma_label"),
            "specialty_raw": clean_spaces(row["specialty"]),
            "school": fields.pop("school"),
        }
        # update_or_create : met à jour si la clé existe, crée sinon
        formation, created = Formation.objects.update_or_create(**natural_key, defaults=fields)
        if created:
            report.created += 1
        else:
            report.updated += 1
        for warning in formation.import_warnings:
            report.warnings.append(f"Ligne {line_number} – {formation} : {warning}")

    return report